
from openpyxl import Workbook
file_name = input('File_name :  ')
sheet_name=input('sheet_name :  ')
wb = openpyxl.load_workbook(file_name, data_only = True)

desengrasantes=[]
for row in range(16, sheet.max_row + 1):
        desengrasantes.append(sheet['C' + str(row)].value)

from openpyxl.chart import BarChart, Reference, Series
values = desengrasantes
chart = BarChart()
chart.add_data(values)
ws.add_chart(chart, "E15")
wb.save("SampleChart.xlsx")
